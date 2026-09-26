from pathlib import Path

from docx import Document

from src.document.reader import DocxReader
from src.document.writer import DocxWriter
from src.models import PIIEntity


INPUT = Path("tests/writer_fixture.docx")
OUTPUT = Path("tests/writer_output.docx")


def make_entity(
    entity_type,
    value,
    start,
    end,
    block_index,
):
    return PIIEntity(
        entity_type=entity_type,
        original_text=value,
        start=start,
        end=end,
        confidence=1.0,
        source="test",
        block_index=block_index,
    )


def test_writer_replaces_pii(tmp_path):

    output = tmp_path / "redacted.docx"

    reader = DocxReader(INPUT)
    blocks = reader.extract_blocks()

    entities = []

    # Find entities manually in fixture blocks.
    for block in blocks:

        if "Sarthak Malvadkar" in block.text:

            start = block.text.index(
                "Sarthak Malvadkar"
            )

            entities.append(
                make_entity(
                    "PERSON",
                    "Sarthak Malvadkar",
                    start,
                    start + len(
                        "Sarthak Malvadkar"
                    ),
                    block.index,
                )
            )

        if (
            "cs.connect@kshinternational.com"
            in block.text
        ):

            start = block.text.index(
                "cs.connect@kshinternational.com"
            )

            entities.append(
                make_entity(
                    "EMAIL_ADDRESS",
                    "cs.connect@kshinternational.com",
                    start,
                    start + len(
                        "cs.connect@kshinternational.com"
                    ),
                    block.index,
                )
            )

        if "+91 81081 14949" in block.text:

            start = block.text.index(
                "+91 81081 14949"
            )

            entities.append(
                make_entity(
                    "PHONE_NUMBER",
                    "+91 81081 14949",
                    start,
                    start + len(
                        "+91 81081 14949"
                    ),
                    block.index,
                )
            )

    replacements = {}

    for entity in entities:

        if entity.entity_type == "PERSON":
            replacement = "Arjun Mehta"

        elif entity.entity_type == "EMAIL_ADDRESS":
            replacement = "contact@example.com"

        elif entity.entity_type == "PHONE_NUMBER":
            replacement = "+91 70000 12345"

        else:
            replacement = "REDACTED"

        replacements[
            (
                entity.block_index,
                entity.start,
                entity.end,
            )
        ] = replacement

    writer = DocxWriter(INPUT)

    writer.save(
        output,
        entities,
        replacements,
    )

    assert output.exists()

    result_document = Document(output)

    full_text = []

    for paragraph in result_document.paragraphs:
        full_text.append(paragraph.text)

    for table in result_document.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text.append(cell.text)

    for section in result_document.sections:

        for paragraph in section.header.paragraphs:
            full_text.append(paragraph.text)

        for paragraph in section.footer.paragraphs:
            full_text.append(paragraph.text)

    result_text = "\n".join(full_text)

    assert "Sarthak Malvadkar" not in result_text
    assert "cs.connect@kshinternational.com" not in result_text
    assert "+91 81081 14949" not in result_text

    assert "Arjun Mehta" in result_text
    assert "contact@example.com" in result_text
    assert "+91 70000 12345" in result_text