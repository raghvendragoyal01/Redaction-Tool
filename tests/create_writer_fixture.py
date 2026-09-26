from pathlib import Path

from docx import Document
from docx.shared import Pt


OUTPUT = Path("tests/writer_fixture.docx")


def main():
    document = Document()

    # Normal paragraph
    paragraph = document.add_paragraph()
    paragraph.add_run("Contact person: ")

    # Bold PII
    bold_run = paragraph.add_run("Sarthak Malvadkar")
    bold_run.bold = True

    paragraph.add_run(" for further information.")

    # PII split across multiple runs
    paragraph = document.add_paragraph()
    paragraph.add_run("Email: ")

    run1 = paragraph.add_run("cs.connect@")
    run1.italic = True

    run2 = paragraph.add_run("kshinternational.com")
    run2.italic = True

    paragraph.add_run(".")

    # Table
    table = document.add_table(
        rows=2,
        cols=2,
    )

    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Sarthak Malvadkar"

    table.cell(1, 0).text = "Email"
    table.cell(1, 1).text = "cs.connect@kshinternational.com"

    # Header
    section = document.sections[0]

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.text = (
        "Confidential: Sarthak Malvadkar"
    )

    # Footer
    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.text = (
        "Contact: +91 81081 14949"
    )

    document.save(OUTPUT)

    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()