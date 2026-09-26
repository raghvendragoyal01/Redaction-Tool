from pathlib import Path

from docx import Document

from src.models import TextBlock


class DocxReader:
    """Read structured text content from a DOCX document."""

    def __init__(self, file_path: str | Path):
        self.file_path = Path(file_path)

        if not self.file_path.exists():
            raise FileNotFoundError(
                f"DOCX file not found: {self.file_path}"
            )

        if self.file_path.suffix.lower() != ".docx":
            raise ValueError(
                f"Expected a DOCX file, got: {self.file_path.suffix}"
            )

        # Load the document exactly once.
        self._document = Document(self.file_path)

    def load(self) -> Document:
        """Return the already-loaded DOCX document."""
        return self._document

    def extract_paragraphs(self) -> list[str]:
        """Extract non-empty body paragraphs."""

        paragraphs = []

        for paragraph in self._document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        return paragraphs

    def extract_tables(self) -> list[list[list[str]]]:
        """Extract text from all tables."""

        tables = []

        for table in self._document.tables:
            rows = []

            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                rows.append(cells)

            tables.append(rows)

        return tables

    @staticmethod
    def iter_paragraphs(doc: Document):
        """
        Unified deterministic generator for all paragraphs in a DOCX document.
        Traverses:
        1. Body paragraphs
        2. Table cells (paragraphs within unique XML cells)
        3. Header paragraphs
        4. Footer paragraphs
        """

        # 1. Body paragraphs
        for p_index, paragraph in enumerate(doc.paragraphs):
            yield f"body.p_{p_index}", "paragraph", paragraph

        # 2. Tables
        for table_index, table in enumerate(doc.tables):
            visited_cells = set()

            for row_index, row in enumerate(table.rows):

                for cell_index, cell in enumerate(row.cells):
                    if cell._tc in visited_cells:
                        continue

                    visited_cells.add(cell._tc)

                    for p_index, paragraph in enumerate(cell.paragraphs):
                        location = (
                            f"table_{table_index}"
                            f".row_{row_index}"
                            f".cell_{cell_index}"
                            f".p_{p_index}"
                        )
                        yield location, "table_cell", paragraph

        # 3. Headers and footers
        for section_index, section in enumerate(doc.sections):

            for p_index, paragraph in enumerate(section.header.paragraphs):
                location = f"section_{section_index}.header.p_{p_index}"
                yield location, "header", paragraph

            for p_index, paragraph in enumerate(section.footer.paragraphs):
                location = f"section_{section_index}.footer.p_{p_index}"
                yield location, "footer", paragraph

    def extract_blocks(self) -> list[TextBlock]:
        """
        Extract document content as structured TextBlock objects.
        """

        blocks: list[TextBlock] = []
        block_index = 0

        for location, block_type, paragraph in self.iter_paragraphs(self._document):
            text = paragraph.text

            if not text.strip():
                continue

            blocks.append(
                TextBlock(
                    text=text,
                    block_type=block_type,
                    location=location,
                    index=block_index,
                    style_name=(
                        paragraph.style.name
                        if paragraph.style
                        else None
                    ),
                )
            )

            block_index += 1

        return blocks

    def extract_text(self) -> str:
        """Extract all document text as plain text."""

        blocks = self.extract_blocks()

        return "\n".join(
            block.text
            for block in blocks
        )